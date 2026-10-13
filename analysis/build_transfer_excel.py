"""地区の格付け・移籍の格上/格下表・回収率変化を Excel にまとめる.
前提: class_transfer.py, track_scale.py, grade_compare.py の出力 (analysis/out/*.csv) があること."""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent / "out"
DEST = Path(__file__).parent / "移籍_格付け表.xlsx"
TIER = {"大井": 1, "船橋": 1, "浦和": 1, "川崎": 1, "兵庫": 2, "高知": 2,
        "愛知": 3, "岩手": 3, "北海道": 3, "金沢": 4, "佐賀": 4, "笠松": 4}
REG = list(TIER)
NAME = {"兵庫": "兵庫(園田・姫路)", "愛知": "愛知(名古屋)", "岩手": "岩手(盛岡・水沢)", "北海道": "北海道(門別)"}
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E78")
SUB = PatternFill("solid", fgColor="DDEBF7")
UP = PatternFill("solid", fgColor="F8CBAD")
DOWN = PatternFill("solid", fgColor="C6EFCE")
SAME = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def body(ws, r0, r1, c0, c1, fmt=None, center=True):
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=F)
            cell.border = BOX
            if center:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if fmt and c > c0 - 1 and isinstance(cell.value, (int, float, str)) and str(cell.value)[:1] in "=-0123456789":
                cell.number_format = fmt


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, bold=True, size=14)
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(name=F, italic=True, color="595959")


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def main():
    scale = pd.read_csv(OUT / "track_scale.csv", encoding="utf-8-sig", index_col=0)
    roi = pd.read_csv(OUT / "transfer_flow_roi.csv", encoding="utf-8-sig").set_index(["元", "先"])
    wb = Workbook()

    # ---- 説明 ----
    ws = wb.active
    ws.title = "説明"
    title(ws, "地方競馬 移籍の格付け表・級の比較")
    lines = [
        ("データ", "地方競馬 2023年1月〜2026年8月の出走馬・払戻データ、2024年1月〜2026年8月のレース一覧"),
        ("格付け", "1着賞金(C級・B級・A級・重賞)を総合して1部〜4部に分類。ばんえいは平地との移籍がないため除外"),
        ("地区", "調教師所属の単位。兵庫=園田+姫路、岩手=盛岡+水沢、愛知=名古屋、北海道=門別。南関東は4場別"),
        ("移籍の判定", "同じ馬(馬名+生年月日)の前走と所属地区が変わった走を移籍とした"),
        ("移籍前", "元の地区での最後の3戦"),
        ("移籍後", "移籍先での1〜3戦目"),
        ("回収率", "全出走に100円ずつ賭けた場合の回収率(%)。確定オッズによる過去実績"),
        ("回収率の変化", "移籍後の回収率 − 移籍前の回収率(ポイント)。地方間移籍の全体平均は単勝+25・複勝+21で、これが基準"),
        ("人気比", "移籍後の複勝率 ÷ 同じ競馬場・同じ人気の馬の平均複勝率。1.00=人気どおり、1.10以上=人気以上、0.90以下=人気以下"),
        ("空欄", "移籍後の出走数が80走未満の組み合わせ(数値が不安定なため表示しない)"),
        ("色", "格上へ=赤系、同格=灰、格下へ=緑系。数値の表は、高いほど緑・低いほど赤のカラースケール"),
        ("注意", "単勝回収率は1頭の高配当で大きく振れるため、組み合わせごとの判断は複勝回収率と人気比を優先する"),
    ]
    for i, (k, v) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=k).font = Font(name=F, bold=True)
        ws.cell(row=i, column=2, value=v).font = Font(name=F)
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, [16, 110])

    # ---- 格付け一覧 ----
    ws = wb.create_sheet("格付け一覧")
    title(ws, "地区の規模ランキング(1着賞金、単位: 万円)", "出典: 2024年1月〜2026年8月のレース一覧。順位はC・B・A級と重賞の賞金を総合して決定")
    cols = ["順位", "格", "地区", "C級", "B級", "A級", "重賞", "全レース中央値", "1着賞金総額(億円/年)", "年間レース数", "在籍頭数(年)"]
    order = REG + ["ばんえい"]
    for j, c in enumerate(cols, 1):
        ws.cell(row=4, column=j, value=c)
    style_header(ws, 4, len(cols))
    for i, reg in enumerate(order, 5):
        s = scale.loc[reg]
        vals = [i - 4 if reg != "ばんえい" else "―", f"{TIER[reg]}部" if reg in TIER else "別枠", NAME.get(reg, reg),
                s["C級中央値"], s["B級中央値"], s["A級中央値"], s["重賞中央値"], s["1着賞金中央値(万円)"],
                round(s["1着賞金総額(億円/年)"], 1), round(s["年間レース数"]), round(s["在籍頭数(年平均)"])]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
        fill = {1: PatternFill("solid", fgColor="FCE4D6"), 2: PatternFill("solid", fgColor="FFF2CC"),
                3: PatternFill("solid", fgColor="E2EFDA"), 4: PatternFill("solid", fgColor="DDEBF7")}.get(TIER.get(reg))
        body(ws, i, i, 1, len(cols))
        if fill:
            for j in (1, 2, 3):
                ws.cell(row=i, column=j).fill = fill
        for j in range(4, 9):
            ws.cell(row=i, column=j).number_format = "#,##0"
        ws.cell(row=i, column=9).number_format = "0.0"
        for j in (10, 11):
            ws.cell(row=i, column=j).number_format = "#,##0"
    widths(ws, [6, 7, 18, 9, 9, 9, 9, 13, 16, 12, 12])
    ws.freeze_panes = "D5"

    # ---- 格上格下表 ----
    ws = wb.create_sheet("格上・同格・格下")
    title(ws, "地区ごとの 格上・ほぼ同格・格下", "実績上の補足: 兵庫は格付け以上に厳しく、大井は南関東の中でも一段上")
    cols = ["地区", "格", "格上", "ほぼ同格", "格下"]
    for j, c in enumerate(cols, 1):
        ws.cell(row=4, column=j, value=c)
    style_header(ws, 4, 5)
    for i, r in enumerate(REG, 5):
        up = "・".join(x for x in REG if TIER[x] < TIER[r]) or "―"
        same = "・".join(x for x in REG if TIER[x] == TIER[r] and x != r) or "―"
        down = "・".join(x for x in REG if TIER[x] > TIER[r]) or "―"
        for j, v in enumerate([NAME.get(r, r), f"{TIER[r]}部", up, same, down], 1):
            ws.cell(row=i, column=j, value=v)
        body(ws, i, i, 1, 5)
        ws.cell(row=i, column=3).fill = UP
        ws.cell(row=i, column=4).fill = SAME
        ws.cell(row=i, column=5).fill = DOWN
        for j in (3, 4, 5):
            ws.cell(row=i, column=j).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    widths(ws, [18, 7, 40, 24, 52])

    # ---- マトリクス ----
    def matrix(sheet, ttl, note, getter, fmt, scale_rule=True, arrows=False):
        ws = wb.create_sheet(sheet)
        title(ws, ttl, note)
        ws.cell(row=4, column=1, value="移籍元 ＼ 移籍先")
        for j, d in enumerate(REG, 2):
            ws.cell(row=4, column=j, value=d)
        style_header(ws, 4, len(REG) + 1)
        for i, r in enumerate(REG, 5):
            c = ws.cell(row=i, column=1, value=r)
            c.font = Font(name=F, bold=True)
            c.fill = SUB
            c.border = BOX
            for j, d in enumerate(REG, 2):
                cell = ws.cell(row=i, column=j)
                cell.value = getter(r, d)
                cell.font = Font(name=F)
                cell.border = BOX
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if fmt:
                    cell.number_format = fmt
                if d == r:
                    cell.fill = PatternFill("solid", fgColor="808080")
                elif arrows:
                    cell.fill = UP if TIER[d] < TIER[r] else (DOWN if TIER[d] > TIER[r] else SAME)
        if scale_rule:
            rng = f"B5:{get_column_letter(len(REG) + 1)}{len(REG) + 4}"
            mid = scale_rule if isinstance(scale_rule, (int, float)) and not isinstance(scale_rule, bool) else None
            ws.conditional_formatting.add(rng, ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="num" if mid is not None else "percentile", mid_value=mid if mid is not None else 50,
                mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
        widths(ws, [16] + [10] * len(REG))
        ws.freeze_panes = "B5"
        return ws

    def arrow(r, d):
        if r == d:
            return None
        k = TIER[r] - TIER[d]
        return "＝ 同格" if k == 0 else ("↑" * min(k, 2) + " 格上へ" if k > 0 else "↓" * min(-k, 2) + " 格下へ")

    def val(col, digits):
        def g(r, d):
            if r == d or (r, d) not in roi.index:
                return None
            return round(float(roi.loc[(r, d), col]), digits)
        return g

    matrix("格の差(早見表)", "移籍元 × 移籍先: 格の差", "↑=格上へ ↑↑=2段以上格上へ ＝=同格 ↓=格下へ ↓↓=2段以上格下へ",
           arrow, None, scale_rule=False, arrows=True)
    matrix("人気比", "移籍後1〜3戦の人気比(複勝率)", "1.00=人気どおり。空欄は80走未満",
           val("移籍後_人気比_複勝率", 2), "0.00", scale_rule=1.0)
    matrix("単勝回収率_移籍後", "移籍後1〜3戦の単勝回収率(%)", "空欄は80走未満。1頭の高配当で大きく振れる点に注意",
           val("移籍後_単勝回収率", 0), "0", scale_rule=70)
    matrix("単勝回収率_変化", "単勝回収率の変化(移籍後 − 移籍前、ポイント)", "全体平均は+25。それより大きければ移籍で好転",
           val("単勝回収率_変化", 0), "+0;-0;0", scale_rule=25)
    matrix("複勝回収率_移籍後", "移籍後1〜3戦の複勝回収率(%)", "空欄は80走未満。90%以上が狙い目の目安",
           val("移籍後_複勝回収率", 0), "0", scale_rule=70)
    matrix("複勝回収率_変化", "複勝回収率の変化(移籍後 − 移籍前、ポイント)", "全体平均は+21",
           val("複勝回収率_変化", 0), "+0;-0;0", scale_rule=21)

    # ---- 方向別集計 (変化は数式) ----
    ws = wb.create_sheet("格の差ごとの集計")
    d = pd.read_csv(OUT / "tier_direction_roi.csv", encoding="utf-8-sig")
    title(ws, "格の差ごとの回収率(移籍前3戦 → 移籍後3戦)", "変化の列は数式(移籍後 − 移籍前)")
    cols = ["方向", "組み合わせ数", "移籍後出走数", "単勝 移籍前(%)", "単勝 移籍後(%)", "単勝 変化",
            "複勝 移籍前(%)", "複勝 移籍後(%)", "複勝 変化", "移籍後の人気比"]
    for j, c in enumerate(cols, 1):
        ws.cell(row=4, column=j, value=c)
    style_header(ws, 4, len(cols))
    for i, row in enumerate(d.itertuples(index=False), 5):
        ws.cell(row=i, column=1, value=row.方向)
        ws.cell(row=i, column=2, value=int(row.組数))
        ws.cell(row=i, column=3, value=int(row.移籍後出走数))
        ws.cell(row=i, column=4, value=float(row.移籍前_単勝))
        ws.cell(row=i, column=5, value=float(row.移籍後_単勝))
        ws.cell(row=i, column=6, value=f"=E{i}-D{i}")
        ws.cell(row=i, column=7, value=float(row.移籍前_複勝))
        ws.cell(row=i, column=8, value=float(row.移籍後_複勝))
        ws.cell(row=i, column=9, value=f"=H{i}-G{i}")
        ws.cell(row=i, column=10, value=float(row.移籍後_人気比))
        body(ws, i, i, 1, len(cols))
        for j in (2, 3):
            ws.cell(row=i, column=j).number_format = "#,##0"
        for j in (4, 5, 7, 8):
            ws.cell(row=i, column=j).number_format = "0"
        for j in (6, 9):
            ws.cell(row=i, column=j).number_format = "+0;-0;0"
        ws.cell(row=i, column=10).number_format = "0.00"
        ws.cell(row=i, column=1).fill = UP if "格上" in row.方向 else (DOWN if "格下" in row.方向 else SAME)
    n = 4 + len(d)
    for col in ("F", "I"):
        ws.conditional_formatting.add(f"{col}5:{col}{n}", ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="max", end_color="63BE7B"))
    widths(ws, [16, 12, 13, 13, 13, 10, 13, 13, 10, 13])

    # ---- 組み合わせ詳細 (変化は数式) ----
    ws = wb.create_sheet("組み合わせ詳細")
    title(ws, "移籍元 → 移籍先 の詳細(移籍後80走以上)", "フィルタ・並べ替えに使えます。変化の列は数式")
    cols = ["移籍元", "移籍先", "格の差", "移籍前 出走数", "移籍前 単勝回収率", "移籍前 複勝回収率", "移籍前 人気比",
            "移籍後 出走数", "移籍後 単勝回収率", "移籍後 複勝回収率", "移籍後 人気比", "単勝 変化", "複勝 変化"]
    for j, c in enumerate(cols, 1):
        ws.cell(row=4, column=j, value=c)
    style_header(ws, 4, len(cols))
    rr = roi.reset_index()
    rr = rr[rr["元"].isin(TIER) & rr["先"].isin(TIER)].copy()
    rr["k"] = rr["元"].map(TIER) - rr["先"].map(TIER)
    rr = rr.sort_values(["k", "移籍後_複勝回収率"], ascending=[False, False])
    i = 4
    for row in rr.itertuples(index=False):
        i += 1
        vals = [row.元, row.先, arrow(row.元, row.先), int(row.移籍前_出走数), round(row.移籍前_単勝回収率, 1),
                round(row.移籍前_複勝回収率, 1), round(row.移籍前_人気比_複勝率, 2), int(row.移籍後_出走数),
                round(row.移籍後_単勝回収率, 1), round(row.移籍後_複勝回収率, 1), round(row.移籍後_人気比_複勝率, 2),
                f"=I{i}-E{i}", f"=J{i}-F{i}"]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
        body(ws, i, i, 1, len(cols))
        ws.cell(row=i, column=3).fill = UP if row.k > 0 else (DOWN if row.k < 0 else SAME)
        for j in (4, 8):
            ws.cell(row=i, column=j).number_format = "#,##0"
        for j in (5, 6, 9, 10):
            ws.cell(row=i, column=j).number_format = "0.0"
        for j in (7, 11):
            ws.cell(row=i, column=j).number_format = "0.00"
        for j in (12, 13):
            ws.cell(row=i, column=j).number_format = "+0.0;-0.0;0.0"
    for col in ("K",):
        ws.conditional_formatting.add(f"{col}5:{col}{i}", ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="num", mid_value=1, mid_color="FFFFFF",
            end_type="max", end_color="63BE7B"))
    for col in ("J",):
        ws.conditional_formatting.add(f"{col}5:{col}{i}", ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="num", mid_value=70, mid_color="FFFFFF",
            end_type="max", end_color="63BE7B"))
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{i}"
    ws.freeze_panes = "C5"
    widths(ws, [10, 10, 12, 11, 13, 13, 11, 11, 13, 13, 11, 10, 10])

    # ---- 級の比較 (grade_compare.py の出力があれば) ----
    def simple_sheet(name, ttl, note, df, fmts=None, fill_col=None, widths_=None, index=False):
        ws = wb.create_sheet(name)
        title(ws, ttl, note)
        df = df.reset_index() if index else df
        for j, c in enumerate(df.columns, 1):
            ws.cell(row=4, column=j, value=str(c))
        style_header(ws, 4, len(df.columns))
        for i, row in enumerate(df.itertuples(index=False), 5):
            for j, v in enumerate(row, 1):
                v = None if (isinstance(v, float) and pd.isna(v)) else v
                if isinstance(v, float):
                    v = round(v, 2)
                ws.cell(row=i, column=j, value=v)
            body(ws, i, i, 1, len(df.columns))
            for j, f in (fmts or {}).items():
                ws.cell(row=i, column=j).number_format = f
        if widths_:
            widths(ws, widths_)
        ws.freeze_panes = "B5"
        return ws

    if (OUT / "grade_band_table.csv").exists():
        band = pd.read_csv(OUT / "grade_band_table.csv", encoding="utf-8-sig").fillna("")
        ws = simple_sheet("級の対応表", "各競馬場の級を1着賞金で横並びにした対応表",
                          "同じ行の級はほぼ同じ賞金水準。重賞・2歳・3歳限定戦は除外。名古屋・笠松の B3〜B12 などは組番号の範囲",
                          band, widths_=[12] + [14] * (len(band.columns) - 1))
        bandfill = ["FCE4D6", "FCE4D6", "FFF2CC", "FFF2CC", "E2EFDA", "E2EFDA", "DDEBF7", "DDEBF7", "EDEDED"]
        for i in range(len(band)):
            ws.cell(row=5 + i, column=1).fill = PatternFill("solid", fgColor=bandfill[i % len(bandfill)])
            for j in range(1, len(band.columns) + 1):
                ws.cell(row=5 + i, column=j).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        lad = pd.read_csv(OUT / "grade_ladder.csv", encoding="utf-8-sig")
        simple_sheet("級別の賞金一覧", "競馬場 × 級 の1着賞金(中央値、万円)", "2024年1月〜2026年8月。15レース以上ある級のみ",
                     lad, fmts={3: "#,##0", 4: "#,##0"}, widths_=[10, 16, 18, 10])
        mv = pd.read_csv(OUT / "grade_move_summary.csv", encoding="utf-8-sig")
        ws = simple_sheet("級の移動と成績", "移籍初戦: 前走の級に対して、移籍先で入った級の賞金水準別の成績",
                          "比 = 移籍先の級の1着賞金 ÷ 前走の級の1着賞金。回収率は%、人気比1.00=人気どおり",
                          mv, fmts={2: "#,##0", 3: "0.0", 4: "0.0", 5: "0.0", 6: "0", 7: "0", 8: "0.00"},
                          widths_=[22, 9, 9, 9, 9, 11, 11, 9])
        pv = pd.read_csv(OUT / "grade_move_by_prev.csv", encoding="utf-8-sig")
        r0 = 5 + len(mv) + 2
        ws.cell(row=r0, column=1, value="複勝率(%): 級の移動 × 前走着順").font = Font(name=F, bold=True)
        for j, c in enumerate(pv.columns, 1):
            ws.cell(row=r0 + 1, column=j, value=str(c))
        style_header(ws, r0 + 1, len(pv.columns))
        for i, row in enumerate(pv.itertuples(index=False), r0 + 2):
            for j, v in enumerate(row, 1):
                ws.cell(row=i, column=j, value=round(v, 1) if isinstance(v, float) else v)
            body(ws, i, i, 1, len(pv.columns))
        ws.conditional_formatting.add(f"B{r0 + 2}:{get_column_letter(len(pv.columns))}{r0 + 1 + len(pv)}", ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFFFFF",
            end_type="max", end_color="63BE7B"))
        nk = pd.read_csv(OUT / "grade_from_nankan_c.csv", encoding="utf-8-sig")
        simple_sheet("南関東C級からの移籍", "南関東のC級から他地区へ移籍した馬の初戦成績",
                     "主な級 = 移籍先で入った級(多い順)。回収率は%", nk,
                     fmts={2: "#,##0", 4: "0", 5: "0.0", 6: "0.0", 7: "0.0", 8: "0", 9: "0"},
                     widths_=[10, 9, 16, 14, 9, 9, 9, 11, 11])
        pr = pd.read_csv(OUT / "grade_pairs.csv", encoding="utf-8-sig")
        ws = simple_sheet("級の対応(実績)", "前走の競馬場・級 → 移籍先の競馬場・級(30件以上)",
                          "実際の移籍で多い級の対応と、移籍初戦の成績", pr,
                          fmts={3: "#,##0", 4: "0.0", 5: "0.00", 6: "0"}, widths_=[14, 14, 8, 9, 9, 11])
        ws.auto_filter.ref = f"A4:F{4 + len(pr)}"

    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(DEST)
    print(DEST)


if __name__ == "__main__":
    main()
